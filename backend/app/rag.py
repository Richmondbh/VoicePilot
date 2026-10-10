"""
Read-only document search (Retrieval-Augmented Generation, the "R" part).

    documents/ (.pdf .docx .txt .md)
        -> extract text (per PDF page, so answers can cite a page)
        -> split into overlapping chunks (~800 characters)
        -> embed every chunk  (Ollama `nomic-embed-text`, local + free)
           or, if Ollama is not available, a TF-IDF vector (scikit-learn)
        -> saved in backend/rag_index.json (rebuilt automatically when files change)
    question -> embed -> cosine similarity -> top-k chunks with their source

This module NEVER executes anything; it only returns text. The agent marks that
text as untrusted before showing it to the LLM.

References:
- Lewis et al. (2020) Retrieval-Augmented Generation: https://arxiv.org/abs/2005.11401
- Ollama embeddings API (/api/embed): https://docs.ollama.com/capabilities/embeddings
- nomic-embed-text task prefixes: https://ollama.com/library/nomic-embed-text
- pypdf text extraction: https://pypdf.readthedocs.io/en/stable/user/extract-text.html
- python-docx: https://python-docx.readthedocs.io/en/latest/
- TF-IDF + cosine similarity: https://scikit-learn.org/stable/modules/feature_extraction.html#tfidf-term-weighting
"""
import json
import math
import re
from pathlib import Path

from . import config

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}
CHUNK_CHARS = 800
CHUNK_OVERLAP = 150
INDEX_VERSION = 2

# Minimum similarity for a chunk to count as relevant. Embedding and TF-IDF
# scores live on different scales, so each has its own threshold. Tune these
# with scripts/evaluate_agent.py if answers are missed or irrelevant text shows up.
MIN_SCORE = {"ollama": 0.45, "tfidf": 0.04}

_index: dict | None = None
_tfidf = None  # (vectorizer, matrix) built in memory for the TF-IDF backend
_pinned = False  # True while tests/evaluation use their own index (set_index)


# ============================================================ loading =====
def _clean(text: str) -> str:
    text = text.replace("\r", "")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _read_pages(path: Path) -> list[tuple[int | None, str]]:
    """Return [(page_number or None, text), ...] for one file."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return [(i + 1, _clean(p.extract_text() or "")) for i, p in enumerate(reader.pages)]
    if suffix == ".docx":
        import docx
        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [(None, _clean("\n".join(parts)))]
    return [(None, _clean(path.read_text(encoding="utf-8", errors="ignore")))]


def chunk_text(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split text into overlapping chunks, preferring to cut at a line or sentence end."""
    chunks, start, n = [], 0, len(text)
    while start < n:
        end = min(n, start + size)
        if end < n:
            cut = max(text.rfind("\n", start + size // 2, end), text.rfind(". ", start + size // 2, end))
            if cut > start:
                end = cut + 1
        piece = text[start:end].strip()
        if len(piece) >= 40:
            chunks.append(piece)
        if end >= n:
            break
        nxt = end - overlap
        space = text.find(" ", nxt)  # don't start in the middle of a word
        start = space + 1 if start < space < end else max(nxt, start + 1)
    return chunks


def _files(folder: Path) -> list[Path]:
    if not folder.exists():
        return []
    return sorted(p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED
                  and not p.name.startswith(("~$", ".")))


def _fingerprint(folder: Path) -> dict:
    return {str(p.relative_to(folder)): [round(p.stat().st_mtime), p.stat().st_size] for p in _files(folder)}


def load_chunks(folder: Path) -> list[dict]:
    chunks = []
    for path in _files(folder):
        try:
            pages = _read_pages(path)
        except Exception as exc:
            print(f"[rag] could not read {path.name}: {exc}")
            continue
        for page, text in pages:
            for piece in chunk_text(text):
                chunks.append({"source": path.name, "path": str(path.relative_to(folder)),
                               "page": page, "text": piece})
    return chunks


# ========================================================= embeddings =====
def _ollama_host() -> str:
    base = config.PROVIDERS["ollama"]["base_url"]
    return re.sub(r"/v1/?$", "", base)


def embed(texts: list[str]) -> list[list[float]]:
    """Embed texts with Ollama. Raises if Ollama or the model is not available."""
    import httpx
    vectors = []
    for i in range(0, len(texts), 32):
        resp = httpx.post(f"{_ollama_host()}/api/embed",
                          json={"model": config.EMBED_MODEL, "input": texts[i:i + 32]}, timeout=180)
        resp.raise_for_status()
        vectors += resp.json()["embeddings"]
    return vectors


def _cosine(a: list[float], b: list[float]) -> float:
    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a))
    nb = math.sqrt(sum(y * y for y in b))
    return dot / (na * nb) if na and nb else 0.0


# ============================================================== index =====
def build_index(folder: Path | None = None, backend: str | None = None) -> dict:
    folder = folder or config.DOCUMENTS_DIR
    backend = backend or config.RAG_BACKEND
    chunks = load_chunks(folder)
    index = {"version": INDEX_VERSION, "folder": str(folder), "fingerprint": _fingerprint(folder),
             "backend": "tfidf", "model": "tfidf", "chunks": chunks, "embeddings": []}
    if chunks and backend in ("auto", "ollama"):
        try:
            # nomic-embed-text expects these task prefixes for documents vs. queries
            index["embeddings"] = embed([f"search_document: {c['text']}" for c in chunks])
            index["backend"], index["model"] = "ollama", config.EMBED_MODEL
        except Exception as exc:
            if backend == "ollama":
                raise
            print(f"[rag] Ollama embeddings unavailable ({exc}); using TF-IDF instead")
    if folder == config.DOCUMENTS_DIR:
        try:
            config.RAG_INDEX_PATH.write_text(json.dumps(index), encoding="utf-8")
        except OSError as exc:
            print(f"[rag] could not save index: {exc}")
    return index


def get_index(force: bool = False) -> dict:
    """Load the saved index, rebuilding it when documents changed (or force=True)."""
    global _index, _tfidf
    if _pinned and _index is not None and not force:
        return _index
    folder = config.DOCUMENTS_DIR
    if not force and _index is not None and _index["fingerprint"] == _fingerprint(folder):
        return _index
    index = None
    if not force and config.RAG_INDEX_PATH.exists():
        try:
            saved = json.loads(config.RAG_INDEX_PATH.read_text(encoding="utf-8"))
            wants_ollama = config.RAG_BACKEND == "ollama" and saved.get("backend") != "ollama"
            if (saved.get("version") == INDEX_VERSION and saved.get("folder") == str(folder)
                    and saved.get("fingerprint") == _fingerprint(folder) and not wants_ollama):
                index = saved
        except (OSError, ValueError):
            pass
    _index, _tfidf = index or build_index(folder), None
    return _index


def set_index(index: dict | None) -> None:
    """Use a specific index (tests / evaluation with another folder). None = back to normal."""
    global _index, _tfidf, _pinned
    _index, _tfidf, _pinned = index, None, index is not None


def _tfidf_scores(index: dict, query: str) -> list[float]:
    """Lexical fallback: average of word TF-IDF and character n-gram TF-IDF.
    Character n-grams let "certification" match "Certified" (no stemming needed)."""
    global _tfidf
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    if _tfidf is None:
        texts = [c["text"] for c in index["chunks"]]
        word = TfidfVectorizer(ngram_range=(1, 2), sublinear_tf=True, stop_words="english")
        char = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5), sublinear_tf=True)
        _tfidf = [(word, word.fit_transform(texts)), (char, char.fit_transform(texts))]
    word, char = [cosine_similarity(vec.transform([query]), matrix)[0] for vec, matrix in _tfidf]
    # A chunk only counts if it shares at least one real word with the question;
    # otherwise character overlap alone ("capital of France") would look relevant.
    return [(w + c) / 2 if w > 0 else 0.0 for w, c in zip(word, char)]


def search(query: str, k: int | None = None, index: dict | None = None) -> list[dict]:
    """Return the k most relevant chunks: [{id, source, page, text, score}, ...]."""
    index = index or get_index()
    if not index["chunks"] or not query.strip():
        return []
    k = k or config.RAG_TOP_K
    if index["backend"] == "ollama":
        try:
            q = embed([f"search_query: {query}"])[0]
            scores = [_cosine(q, e) for e in index["embeddings"]]
        except Exception as exc:  # Ollama stopped after indexing -> degrade gracefully
            print(f"[rag] query embedding failed ({exc}); using TF-IDF for this search")
            scores, index = _tfidf_scores(index, query), {**index, "backend": "tfidf"}
    else:
        scores = _tfidf_scores(index, query)
    threshold = MIN_SCORE[index["backend"]]
    ranked = sorted(range(len(scores)), key=lambda i: -scores[i])[:k]
    hits = []
    for i in ranked:
        if scores[i] < threshold:
            continue
        c = index["chunks"][i]
        hits.append({"id": f"S{len(hits) + 1}", "source": c["source"], "page": c["page"],
                     "text": c["text"], "score": round(float(scores[i]), 3)})
    return hits


def source_label(hit: dict) -> str:
    return f"{hit['source']} p.{hit['page']}" if hit.get("page") else hit["source"]


def list_documents() -> dict:
    index = get_index()
    counts: dict[str, int] = {}
    for c in index["chunks"]:
        counts[c["path"]] = counts.get(c["path"], 0) + 1
    files = [{"path": p, "chunks": counts.get(p, 0)} for p in index["fingerprint"]]
    return {"folder": index["folder"], "backend": index["backend"], "model": index["model"],
            "files": files, "chunks": len(index["chunks"])}
